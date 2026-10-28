from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
from pathlib import Path

src=Path(r'C:\Users\23716\Desktop\test\test111\面试复习讲义_生成器装饰器迭代器.docx')
out=Path(r'C:\Users\23716\Desktop\test\test111\面试复习讲义_生成器装饰器迭代器.pdf')
font=r'C:\Windows\Fonts\msyh.ttc'
pdfmetrics.registerFont(TTFont('YaHei',font, subfontIndex=0))
pdfmetrics.registerFont(TTFont('YaHeiBold',font, subfontIndex=0))
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyCN', parent=styles['BodyText'], fontName='YaHei', fontSize=8.8, leading=11, spaceAfter=3))
styles.add(ParagraphStyle(name='TitleCN', parent=styles['Title'], fontName='YaHei', fontSize=17, leading=21, alignment=TA_CENTER, textColor=colors.black, spaceAfter=5))
styles.add(ParagraphStyle(name='SubCN', parent=styles['BodyText'], fontName='YaHei', fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.HexColor('#555555'), spaceAfter=8))
styles.add(ParagraphStyle(name='H1CN', parent=styles['Heading1'], fontName='YaHei', fontSize=12, leading=15, spaceBefore=7, spaceAfter=3, textColor=colors.black, keepWithNext=True))
styles.add(ParagraphStyle(name='H2CN', parent=styles['Heading2'], fontName='YaHei', fontSize=9.8, leading=12, spaceBefore=5, spaceAfter=2, textColor=colors.black, keepWithNext=True))
styles.add(ParagraphStyle(name='CodeCN', parent=styles['Code'], fontName='Courier', fontSize=7.3, leading=8.7, leftIndent=8, rightIndent=4, spaceAfter=1))
styles.add(ParagraphStyle(name='QCN', parent=styles['BodyText'], fontName='YaHei', fontSize=8.6, leading=10.5, spaceAfter=1))

def clean(s): return escape(s).replace('\\n','<br/>')
def p(text, st='BodyCN'): return Paragraph(clean(text), styles[st])

docx=Document(str(src)); story=[]
for para in docx.paragraphs:
    txt=para.text
    if not txt: continue
    sty=para.style.name
    if sty=='Title': story.append(p(txt,'TitleCN'))
    elif sty=='Heading 1': story.append(p(txt,'H1CN'))
    elif sty=='Heading 2': story.append(p(txt,'H2CN'))
    elif sty=='Code': story.append(p(txt,'CodeCN'))
    elif sty=='List Bullet': story.append(p('• '+txt,'BodyCN'))
    else: story.append(p(txt,'BodyCN'))
    
for t in docx.tables:
    data=[]
    for ri,row in enumerate(t.rows):
        vals=[]
        for cell in row.cells:
            vals.append(Paragraph(clean(cell.text), styles['BodyCN']))
        data.append(vals)
    widths=None
    if len(data[0])==3: widths=[28*mm,83*mm,51*mm]
    elif len(data[0])==4: widths=[38*mm,28*mm,31*mm,55*mm]
    else: widths=[35*mm]*len(data[0])
    tab=Table(data,colWidths=widths,repeatRows=1,hAlign='CENTER')
    ts=[('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#D9D9D9')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F4E78')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]
    for i in range(1,len(data)):
        if i%2==0: ts.append(('BACKGROUND',(0,i),(-1,i),colors.HexColor('#F3F6F9')))
    tab.setStyle(TableStyle(ts)); story.append(Spacer(1,2)); story.append(tab); story.append(Spacer(1,4))

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont('YaHei',7); canvas.setFillColor(colors.HexColor('#777777')); canvas.drawCentredString(A4[0]/2,10*mm,'生成器 装饰器 迭代器面试复习讲义    第 %d 页'%doc.page); canvas.restoreState()

doc=SimpleDocTemplate(str(out),pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=12*mm,bottomMargin=15*mm)
doc.build(story,onFirstPage=footer,onLaterPages=footer)
print(out)
